from pathlib import Path
from textwrap import wrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = ROOT / "assets" / "images"
DOCS.mkdir(exist_ok=True)
OUTPUT = DOCS / "Laporan_Tahap_2_dan_3_TEMU_UIN_Malang.docx"


def font(size=24, bold=False):
    candidates = [
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeui.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def make_diagram(path, title, boxes, arrows, width=1500, height=720):
    image = Image.new("RGB", (width, height), "#F7F5EE")
    draw = ImageDraw.Draw(image)
    draw.text((50, 28), title, fill="#053B2B", font=font(34, True))
    for box in boxes:
        x, y, w, h, label, color = box
        draw.rounded_rectangle((x, y, x + w, y + h), radius=22,
                               fill=color, outline="#006B45", width=4)
        lines = wrap(label, width=max(12, int(w / 18)))
        total = len(lines) * 30
        current_y = y + (h - total) / 2
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=font(23, True))
            text_w = bbox[2] - bbox[0]
            draw.text((x + (w - text_w) / 2, current_y), line,
                      fill="#14231D", font=font(23, True))
            current_y += 30
    for x1, y1, x2, y2 in arrows:
        draw.line((x1, y1, x2, y2), fill="#C9A227", width=6)
        draw.polygon([(x2, y2), (x2 - 16, y2 - 10), (x2 - 16, y2 + 10)],
                     fill="#C9A227")
    image.save(path)


def create_diagrams():
    flow_boxes = [
        (70, 200, 210, 100, "Mulai / Login", "#E8F3ED"),
        (340, 200, 230, 100, "Validasi akun", "#FFF3D4"),
        (640, 200, 230, 100, "Dashboard", "#E8F3ED"),
        (950, 90, 230, 100, "Jelajah laporan", "#FFFFFF"),
        (950, 230, 230, 100, "Buat laporan", "#FFFFFF"),
        (950, 370, 230, 100, "Chat & notifikasi", "#FFFFFF"),
        (1250, 200, 180, 100, "Selesai / Logout", "#FFF3D4"),
    ]
    flow_arrows = [
        (280, 250, 340, 250), (570, 250, 640, 250),
        (870, 230, 950, 140), (870, 250, 950, 280),
        (870, 270, 950, 420), (1180, 140, 1250, 230),
        (1180, 280, 1250, 250), (1180, 420, 1250, 270),
    ]
    make_diagram(DOCS / "diagram_alur_aplikasi.png",
                 "Diagram Alur Utama TEMU UIN Malang", flow_boxes, flow_arrows)

    nav_boxes = [
        (80, 220, 260, 110, "Login Mahasiswa", "#E8F3ED"),
        (80, 430, 260, 110, "Login Admin", "#FFF3D4"),
        (500, 160, 300, 110, "Dashboard Mahasiswa", "#E8F3ED"),
        (500, 390, 300, 110, "Dashboard Admin", "#FFF3D4"),
        (980, 70, 220, 90, "Jelajah", "#FFFFFF"),
        (980, 180, 220, 90, "Lapor", "#FFFFFF"),
        (980, 290, 220, 90, "Chat", "#FFFFFF"),
        (980, 400, 220, 90, "Profil", "#FFFFFF"),
        (1300, 180, 180, 90, "Laporan", "#FFFFFF"),
        (1300, 300, 180, 90, "Chat", "#FFFFFF"),
        (1300, 420, 180, 90, "Notifikasi", "#FFFFFF"),
    ]
    nav_arrows = [
        (340, 275, 500, 215), (340, 485, 500, 445),
        (800, 215, 980, 115), (800, 215, 980, 225),
        (800, 215, 980, 335), (800, 215, 980, 445),
        (800, 445, 1300, 225), (800, 445, 1300, 345),
        (800, 445, 1300, 465),
    ]
    make_diagram(DOCS / "struktur_navigasi.png",
                 "Struktur Navigasi Aplikasi", nav_boxes, nav_arrows,
                 height=650)

    data_boxes = [
        (80, 150, 290, 220, "users\n(id, nim, email, nama)", "#E8F3ED"),
        (510, 80, 320, 240, "reports\n(id, nama, lokasi, jenis,\npelapor, status, foto)", "#FFF3D4"),
        (980, 80, 330, 240, "chats\n(id, barang_id, owner_email,\nis_online, unread_count)", "#E8F3ED"),
        (980, 420, 330, 180, "messages\n(chat_id, pengirim,\nteks, waktu)", "#FFFFFF"),
        (510, 440, 320, 160, "notifications\n(email, title, message,\nwaktu, is_read)", "#FFFFFF"),
        (80, 450, 290, 150, "admins\n(config backend + token)", "#FFF3D4"),
    ]
    data_arrows = [
        (370, 230, 510, 200), (830, 200, 980, 200),
        (1145, 320, 1145, 420), (830, 500, 510, 520),
        (370, 510, 510, 520), (370, 245, 980, 500),
    ]
    make_diagram(DOCS / "model_data.png",
                 "Rancangan Model Data", data_boxes, data_arrows,
                 height=700)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9E5DE", size="8"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="14231D", size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_table(doc, headers, rows, widths=None, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True, "FFFFFF", font_size)
        set_cell_shading(table.rows[0].cells[i], "006B45")
        set_cell_border(table.rows[0].cells[i], "006B45")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, "14231D", font_size)
            set_cell_shading(cells[i], "F7F5EE" if len(table.rows) % 2 == 0 else "FFFFFF")
            set_cell_border(cells[i])
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(111, 125, 118)


def add_placeholder(doc, title, instruction, height=2.7):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F7F4")
    set_cell_border(cell, "8BB9A2", "14")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run("SCREENSHOT BUKTI\n" + title)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0, 107, 69)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(12)
    r2 = p2.add_run(instruction)
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(111, 125, 118)
    for _ in range(max(1, int(height / 0.45))):
        cell.add_paragraph("")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    set_cell_border(cell, "D0D0D0")
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(20, 35, 29)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color in [
        ("Title", 23, "053B2B"),
        ("Heading 1", 16, "006B45"),
        ("Heading 2", 13, "053B2B"),
        ("Heading 3", 11, "006B45"),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("TEMU UIN MALANG  |  LAPORAN TAHAP 2 DAN 3")
    hr.font.size = Pt(8)
    hr.font.color.rgb = RGBColor(111, 125, 118)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Dokumen laporan akademik - ")
    fr.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)


def build_report():
    create_diagrams()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Laporan Tahap 2 dan Tahap 3 TEMU UIN Malang"
    doc.core_properties.subject = "Perancangan Sistem dan Implementasi Aplikasi Flutter"
    doc.core_properties.author = "[Nama Mahasiswa]"
    doc.core_properties.keywords = "Flutter, UIN Malang, Lost and Found, TEMU"

    logo = ASSETS / "logo_uin_transparent.png"
    if logo.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(logo), width=Cm(3.0))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LAPORAN PERKEMBANGAN PROJECT")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(5, 59, 43)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TAHAP 2: PERANCANGAN SISTEM DAN ANTARMUKA\n"
                  "TAHAP 3: IMPLEMENTASI APLIKASI DENGAN FLUTTER")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0, 107, 69)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TEMU - Lost & Found UIN Malang")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(201, 162, 39)
    doc.add_paragraph()
    add_table(doc, ["Identitas", "Keterangan"], [
        ("Nama mahasiswa", "[Isi nama mahasiswa]"),
        ("NIM", "[Isi NIM]"),
        ("Kelas / mata kuliah", "[Isi kelas dan mata kuliah]"),
        ("Dosen pengampu", "[Isi nama dosen]"),
        ("Platform", "Flutter: Web, Android, iOS, Windows, macOS, Linux"),
        ("Tanggal laporan", "7 Oktober 2026"),
    ], widths=[5, 10], font_size=9)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Dokumen ini menjelaskan rancangan dan implementasi MVP aplikasi TEMU UIN Malang.")
    r.italic = True
    r.font.color.rgb = RGBColor(111, 125, 118)
    doc.add_page_break()

    doc.add_heading("Ringkasan Project", level=1)
    doc.add_paragraph(
        "TEMU UIN Malang adalah aplikasi Lost & Found kampus yang membantu mahasiswa "
        "melaporkan barang hilang atau ditemukan, mencari laporan, berkomunikasi melalui "
        "chat, dan memantau proses klaim. Aplikasi menggunakan Flutter untuk antarmuka "
        "lintas platform, REST API PHP sebagai layanan backend, MySQL/MariaDB sebagai "
        "database utama, dan Hive sebagai cache atau fallback lokal."
    )
    doc.add_paragraph(
        "MVP yang dihasilkan telah mencakup autentikasi mahasiswa berbasis email resmi "
        "UIN Malang, login admin terpisah, dashboard, laporan, upload foto, chat, "
        "notifikasi popup, badge jumlah pesan, status online/last seen, database, "
        "validasi masukan, pengujian, serta repository Git."
    )
    doc.add_heading("Tautan Bukti Project", level=2)
    p = doc.add_paragraph("Repository GitHub: ")
    add_hyperlink(p, "mustikaprmataputrii/lost_and_found_campus",
                  "https://github.com/mustikaprmataputrii/lost_and_found_campus")
    p = doc.add_paragraph("Dokumentasi README pada repository: ")
    add_hyperlink(p, "README.md",
                  "https://github.com/mustikaprmataputrii/lost_and_found_campus/blob/main/README.md")
    p = doc.add_paragraph("Backend API pada repository: ")
    add_hyperlink(p, "backend/public/api/index.php",
                  "https://github.com/mustikaprmataputrii/lost_and_found_campus/blob/main/backend/public/api/index.php")
    doc.add_page_break()

    doc.add_heading("BAB I - PENDAHULUAN", level=1)
    doc.add_heading("1.1 Latar Belakang", level=2)
    doc.add_paragraph(
        "Informasi barang hilang atau ditemukan di lingkungan kampus sering tersebar "
        "di grup percakapan dan sulit ditelusuri. TEMU dirancang sebagai ruang terpusat "
        "untuk menyimpan laporan, mempertemukan pelapor dengan pihak yang mencari, dan "
        "mendukung proses komunikasi secara lebih terstruktur."
    )
    doc.add_heading("1.2 Tujuan", level=2)
    add_bullets(doc, [
        "Menyediakan media pelaporan barang hilang dan ditemukan di lingkungan UIN Malang.",
        "Membantu mahasiswa menjelajah laporan berdasarkan nama, deskripsi, jenis, dan lokasi.",
        "Menyediakan chat dan notifikasi untuk mendukung proses klaim.",
        "Menyediakan dashboard admin untuk memantau seluruh laporan, chat, dan notifikasi.",
        "Menerapkan aplikasi Flutter secara terstruktur dengan pemisahan UI, logika, data, dan layanan.",
    ])
    doc.add_heading("1.3 Sasaran Pengguna", level=2)
    add_table(doc, ["Aktor", "Karakteristik", "Kebutuhan"], [
        ("Mahasiswa / civitas", "Memiliki NIM dan email resmi kampus",
         "Membuat laporan, mencari barang, chat, dan melihat notifikasi"),
        ("Admin", "Pengelola aplikasi",
         "Melihat seluruh data laporan, chat, notifikasi, dan memantau aktivitas"),
    ], widths=[4, 5, 7], font_size=8)
    doc.add_page_break()

    doc.add_heading("BAB II - TAHAP 2: PERANCANGAN SISTEM DAN ANTARMUKA", level=1)
    doc.add_heading("2.1 Analisis Kebutuhan Fungsional", level=2)
    add_table(doc, ["Kode", "Kebutuhan", "Aktor"], [
        ("F-01", "Login mahasiswa dengan format NIM@student.uin-malang.ac.id", "Mahasiswa"),
        ("F-02", "Login admin dengan email dan password yang diverifikasi backend", "Admin"),
        ("F-03", "Melihat dashboard dan ringkasan aktivitas", "Mahasiswa/Admin"),
        ("F-04", "Menjelajah, mencari, dan memfilter laporan", "Mahasiswa"),
        ("F-05", "Membuat laporan barang hilang atau ditemukan", "Mahasiswa"),
        ("F-06", "Memilih foto dari kamera atau galeri", "Mahasiswa"),
        ("F-07", "Membuka chat, mengirim pesan, dan melihat online/last seen", "Mahasiswa/Admin"),
        ("F-08", "Menerima popup notifikasi dan badge jumlah pesan", "Mahasiswa"),
        ("F-09", "Melihat seluruh data melalui dashboard admin", "Admin"),
    ], widths=[2, 10, 3], font_size=7)
    doc.add_heading("2.2 Kebutuhan Nonfungsional", level=2)
    add_bullets(doc, [
        "Kemudahan penggunaan: alur login, dashboard, laporan, dan chat dibuat singkat.",
        "Konsistensi: warna hijau UIN, emas, dan sand digunakan di seluruh halaman.",
        "Keterbacaan: ukuran teks, kontras, label, status, dan validasi dibuat jelas.",
        "Lintas platform: target Flutter mencakup web, Android, iOS, Windows, macOS, dan Linux.",
        "Keandalan data: MySQL/MariaDB menjadi penyimpanan utama dan Hive menjadi fallback lokal.",
    ])
    doc.add_heading("2.3 Diagram Alur Aplikasi", level=2)
    doc.add_picture(str(DOCS / "diagram_alur_aplikasi.png"), width=Cm(16))
    add_caption(doc, "Gambar 1. Diagram alur utama aplikasi TEMU")
    doc.add_heading("2.4 Kasus Penggunaan dan Skenario", level=2)
    add_table(doc, ["Kode", "Kasus penggunaan", "Skenario ringkas", "Hasil"], [
        ("UC-01", "Login mahasiswa", "Mahasiswa mengisi nama dan email kampus yang valid.", "Masuk ke dashboard mahasiswa"),
        ("UC-02", "Login admin", "Admin memilih mode admin lalu memasukkan email dan password.", "Masuk ke dashboard admin"),
        ("UC-03", "Mencari laporan", "Pengguna membuka Jelajah, mencari nama/deskripsi, dan memakai filter.", "Daftar laporan sesuai kriteria"),
        ("UC-04", "Membuat laporan", "Pengguna mengisi nama barang, lokasi, detail, kontak, jenis, dan foto.", "Laporan divalidasi dan tersimpan"),
        ("UC-05", "Chat klaim", "Pengguna membuka detail barang dan mengirim pesan.", "Pesan masuk ke sesi chat"),
        ("UC-06", "Notifikasi", "Pesan baru memicu popup dan badge jumlah.", "Pengguna mengetahui aktivitas baru"),
        ("UC-07", "Monitoring admin", "Admin membuka tab laporan, chat, atau notifikasi.", "Admin melihat seluruh data"),
    ], widths=[2, 3, 8, 3], font_size=7)
    doc.add_heading("2.5 Struktur Navigasi Aplikasi", level=2)
    doc.add_picture(str(DOCS / "struktur_navigasi.png"), width=Cm(16))
    add_caption(doc, "Gambar 2. Struktur navigasi mahasiswa dan admin")
    doc.add_heading("2.6 Rancangan Basis Data / Model Data", level=2)
    doc.add_picture(str(DOCS / "model_data.png"), width=Cm(16))
    add_caption(doc, "Gambar 3. Model data yang digunakan")
    add_table(doc, ["Entitas", "Atribut penting", "Fungsi"], [
        ("users", "id, nim, email, nama", "Menyimpan identitas mahasiswa yang login"),
        ("reports", "id, nama, lokasi, jenis, pelapor, status, foto", "Menyimpan laporan barang"),
        ("chats", "id, barang_id, owner_email, is_online, unread_count", "Menyimpan sesi percakapan"),
        ("messages", "chat_id, pengirim, teks, waktu, is_me", "Menyimpan isi pesan"),
        ("notifications", "email, title, message, waktu, is_read", "Menyimpan notifikasi pengguna"),
        ("admin config", "email, password, name, secret", "Konfigurasi akses admin backend"),
    ], widths=[3, 8, 6], font_size=8)
    doc.add_heading("2.7 Sketsa Awal / Wireframe", level=2)
    doc.add_paragraph("Sketsa berikut menggambarkan struktur halaman sebelum implementasi visual.")
    add_table(doc, ["Halaman", "Struktur wireframe"], [
        ("Login", "Logo UIN - nama aplikasi - input nama/email atau email/password admin - tombol masuk - pilihan mode"),
        ("Dashboard mahasiswa", "AppBar - salam pengguna - kartu ringkasan - akses cepat Jelajah/Lapor - ringkasan aktivitas"),
        ("Jelajah", "AppBar - search - filter - daftar kartu barang - status ditemukan/hilang - lokasi"),
        ("Lapor", "Form nama, lokasi, detail, kontak, jenis - kartu foto - tombol Kamera/Galeri - terbitkan"),
        ("Chat", "Daftar sesi - preview pesan - badge unread - Online/last seen - room chat"),
        ("Profil", "Logo - identitas - status verifikasi - waktu sesi - tombol logout"),
        ("Admin", "AppBar admin - kartu statistik - tab Laporan, Chat, Notifikasi - detail data"),
    ], widths=[4, 13], font_size=8)
    doc.add_heading("2.8 Purwarupa Antarmuka", level=2)
    doc.add_paragraph(
        "Purwarupa dikembangkan dengan nuansa UIN Malang melalui warna hijau tua, hijau "
        "utama, emas, dan sand. Logo UIN digunakan sebagai identitas visual pada login, "
        "dashboard, AppBar, dan halaman admin. Komponen Material 3 dipakai agar tampilan "
        "konsisten pada perangkat berbeda."
    )
    add_placeholder(doc, "Purwarupa login dan dashboard",
                    "Tempel screenshot halaman login dan dashboard mahasiswa dari Chrome/Android.", 2.0)
    add_placeholder(doc, "Purwarupa laporan dan chat",
                    "Tempel screenshot form laporan, chat, badge pesan, serta status online.", 2.0)
    doc.add_page_break()

    doc.add_heading("BAB III - TAHAP 3: IMPLEMENTASI APLIKASI DENGAN FLUTTER", level=1)
    doc.add_heading("3.1 Teknologi yang Digunakan", level=2)
    add_table(doc, ["Komponen", "Teknologi / Implementasi"], [
        ("Frontend", "Dart dan Flutter Material 3"),
        ("Database utama", "MySQL/MariaDB melalui Laragon"),
        ("Backend", "PHP REST API"),
        ("Cache lokal", "Hive CE dan SharedPreferences"),
        ("Upload foto", "image_picker; kamera/galeri pada perangkat yang didukung"),
        ("HTTP", "Package http untuk komunikasi REST API"),
        ("Pengujian", "Flutter widget test dan analyzer"),
        ("Repository", "GitHub"),
    ], widths=[5, 12], font_size=8)
    doc.add_heading("3.2 Struktur Kode Terstruktur", level=2)
    add_code(doc, """lib/
  main.dart                    Entry point dan registrasi part
  core/theme/                  Tema UIN dan Material 3
  core/utils/                  Formatter tanggal dan status
  models/                      Barang, chat, pesan, notifikasi
  repositories/                Local, mock, dan remote data
  services/                    Auth, API, session
  screens/                     Login, dashboard, laporan, chat, admin
  widgets/                     Komponen UI bersama

backend/
  database/schema.sql          Struktur MySQL
  public/api/index.php         REST API
  config.php                   Konfigurasi database dan admin""")
    doc.add_heading("3.3 Implementasi Navigasi", level=2)
    add_bullets(doc, [
        "Login mahasiswa berhasil menuju DashboardScreen lalu MainNavigationScreen.",
        "MainNavigationScreen menggunakan NavigationBar untuk Jelajah, Lapor, Chat, dan Profil.",
        "DetailBarangScreen membuka RoomChatScreen untuk proses klaim.",
        "Login admin berhasil menuju AdminDashboardScreen yang memiliki tab Laporan, Chat, dan Notifikasi.",
        "Logout membersihkan session dan mengembalikan pengguna ke LoginScreen.",
    ])
    doc.add_heading("3.4 Pengelolaan Masukan dan Validasi", level=2)
    add_table(doc, ["Masukan", "Validasi"], [
        ("Email mahasiswa", "Wajib mengikuti pola angka@student.uin-malang.ac.id"),
        ("Nama mahasiswa", "Tidak boleh kosong"),
        ("Email/password admin", "Wajib diisi dan diverifikasi backend"),
        ("Form laporan", "Nama, lokasi, deskripsi, dan kontak wajib diisi"),
        ("Pesan chat", "Pesan kosong tidak dapat dikirim"),
        ("Foto", "Dapat dipilih dari kamera atau galeri sesuai dukungan platform"),
    ], widths=[6, 11], font_size=8)
    doc.add_heading("3.5 Pengelolaan Data", level=2)
    doc.add_paragraph(
        "AppDataRepository menjadi penghubung antara UI, API, dan penyimpanan lokal. "
        "Saat API tersedia, data disimpan dan dimuat dari MySQL melalui endpoint REST. "
        "Saat API tidak tersedia, aplikasi menggunakan cache Hive agar MVP masih dapat "
        "dibuka dan diuji. Endpoint admin memakai token sesi untuk memuat seluruh data."
    )
    add_code(doc, """POST /api/index.php?path=auth/login
POST /api/index.php?path=auth/admin-login
GET  /api/index.php?path=sync&email=...
GET  /api/index.php?path=admin/sync
POST /api/index.php?path=sync""")
    doc.add_heading("3.6 Implementasi Fitur Utama", level=2)
    add_bullets(doc, [
        "Login mahasiswa dengan pembatasan domain email resmi kampus.",
        "Login admin dan dashboard monitoring terpisah dari dashboard mahasiswa.",
        "Laporan barang hilang/ditemukan dengan status dan detail.",
        "Upload foto kamera atau galeri melalui image_picker.",
        "Chat klaim dengan status Online atau Terakhir dilihat.",
        "Popup notifikasi pesan dan badge jumlah chat/pemberitahuan.",
        "Database MySQL/MariaDB dan fallback Hive lokal.",
    ])
    doc.add_heading("3.7 Pengujian dan Hasil Verifikasi", level=2)
    add_table(doc, ["Pemeriksaan", "Hasil", "Keterangan"], [
        ("Flutter test", "LULUS", "10 test login, validasi, navigasi, laporan, foto, dan chat"),
        ("Flutter analyze", "LULUS", "No issues found"),
        ("PHP lint", "LULUS", "Tidak ada syntax error pada API dan config"),
        ("API admin login", "LULUS", "Menghasilkan token admin"),
        ("API admin sync", "LULUS", "Mengembalikan seluruh laporan, chat, dan notifikasi"),
        ("Flutter build web", "LULUS", "Build Web berhasil dikompilasi"),
        ("Git", "LULUS", "Commit tersedia pada branch main"),
    ], widths=[5, 3, 9], font_size=8)
    add_placeholder(doc, "Bukti terminal testing",
                    "Tempel screenshot terminal yang menampilkan All tests passed, No issues found, dan PHP lint.", 1.7)
    doc.add_heading("3.8 Repository dan Dokumentasi", level=2)
    doc.add_paragraph(
        "Kode sumber dan dokumentasi project telah diunggah ke repository GitHub. "
        "Repository berisi source Flutter, backend, database schema, test, README, dan "
        "konfigurasi platform."
    )
    p = doc.add_paragraph("Link repository: ")
    add_hyperlink(p, "https://github.com/mustikaprmataputrii/lost_and_found_campus",
                  "https://github.com/mustikaprmataputrii/lost_and_found_campus")
    p = doc.add_paragraph("Commit utama implementasi admin: 732d5fa")
    p.runs[0].font.name = "Consolas"
    doc.add_page_break()

    doc.add_heading("BAB IV - BUKTI YANG PERLU DILAMPIRKAN", level=1)
    doc.add_paragraph(
        "Bagian ini berisi daftar screenshot yang perlu diambil oleh mahasiswa. "
        "Placeholder pada dokumen dapat diganti dengan gambar asli sebelum dikirim kepada dosen."
    )
    add_table(doc, ["No.", "Screenshot yang diperlukan", "Cara mengambil bukti"], [
        ("1", "Login mahasiswa", "Tampilkan halaman login dengan logo UIN dan validasi email kampus."),
        ("2", "Dashboard mahasiswa", "Login sebagai mahasiswa lalu screenshot halaman Dashboard."),
        ("3", "Jelajah laporan", "Tampilkan daftar laporan, pencarian, filter, dan kartu barang."),
        ("4", "Form laporan", "Tampilkan form validasi dan bottom sheet pilihan Kamera/Galeri."),
        ("5", "Kamera perangkat", "Ambil screenshot kamera langsung pada Android/iOS; Chrome laptop dapat membuka file picker."),
        ("6", "Chat", "Tampilkan daftar chat, badge angka, room chat, dan Online/last seen."),
        ("7", "Notifikasi", "Tampilkan popup pesan masuk dan halaman pemberitahuan."),
        ("8", "Dashboard admin", "Login admin lalu screenshot statistik dan tab Laporan/Chat/Notifikasi."),
        ("9", "Database", "Buka phpMyAdmin database lost_and_found_campus dan screenshot tabel users, reports, chats, messages, notifications."),
        ("10", "Testing", "Screenshot terminal hasil flutter test, analyze, build web, dan PHP lint."),
        ("11", "GitHub", "Buka repository GitHub dan screenshot daftar file serta commit terbaru."),
    ], widths=[1.2, 6.5, 9.3], font_size=7)
    doc.add_heading("Catatan Bukti Kamera", level=2)
    doc.add_paragraph(
        "Pada Android dan iOS, ImageSource.camera memanggil kamera perangkat. Pada Chrome "
        "desktop, browser dapat menampilkan dialog pemilihan file karena perilaku input "
        "kamera web bergantung pada dukungan browser. Oleh karena itu, bukti kamera langsung "
        "sebaiknya diambil dari Android atau iOS."
    )
    add_placeholder(doc, "Database MySQL / phpMyAdmin",
                    "Tempel screenshot phpMyAdmin: database lost_and_found_campus dan daftar tabel.", 2.0)
    add_placeholder(doc, "Repository GitHub",
                    "Tempel screenshot repository GitHub yang menampilkan file project dan branch main.", 2.0)
    doc.add_page_break()

    doc.add_heading("BAB V - KESIMPULAN", level=1)
    doc.add_paragraph(
        "Tahap 2 telah menghasilkan rancangan alur aplikasi, use case, struktur navigasi, "
        "model data, wireframe, dan purwarupa antarmuka. Tahap 3 telah mengimplementasikan "
        "rancangan tersebut menjadi MVP Flutter yang dapat dijalankan pada target platform "
        "yang tersedia, terhubung ke backend REST API dan database MySQL/MariaDB, serta "
        "memiliki validasi masukan dan pengujian."
    )
    doc.add_paragraph(
        "Dengan adanya dashboard mahasiswa dan dashboard admin, aplikasi telah memiliki "
        "pemisahan akses sesuai peran pengguna. Fitur utama Lost & Found—laporan, pencarian, "
        "chat, notifikasi, upload foto, penyimpanan data, dan monitoring admin—telah dapat "
        "digunakan sebagai dasar pengembangan tahap berikutnya."
    )
    doc.add_heading("Pengembangan Lanjutan", level=2)
    add_bullets(doc, [
        "Mengganti kredensial admin demo dengan password hash dan manajemen akun admin pada database.",
        "Menyempurnakan kamera web desktop menggunakan preview webcam berbasis getUserMedia.",
        "Menambahkan autentikasi token yang lebih kuat, role permission, dan audit log.",
        "Menambahkan notifikasi push dan deployment backend ke server publik.",
    ])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()

