from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Laporan_Tahap_4_dan_5_TEMU_UIN_Malang.docx"
SCREENSHOT_LOGIN_BEFORE = Path(
    r"C:\Users\L E N O V O\OneDrive\Pictures\Screenshots\Screenshot 2026-10-09 141158.png"
)
SCREENSHOT_ADMIN_BEFORE = Path(
    r"C:\Users\L E N O V O\OneDrive\Pictures\Screenshots\Screenshot 2026-10-09 134021.png"
)
GREEN = "006B45"
DARK = "053B2B"
PALE = "F2F7F4"
INK = "14231D"


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    props.append(element)


def set_cell(cell, value, *, bold=False, color=INK, size=8):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(str(value))
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)


def table(doc, headers, rows, size=8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(headers):
        set_cell(t.rows[0].cells[i], text, bold=True, color="FFFFFF", size=size)
        shade(t.rows[0].cells[i], GREEN)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            set_cell(cells[i], text, size=size)
            shade(cells[i], "FFFFFF" if len(t.rows) % 2 else PALE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def bullets(doc, values):
    for value in values:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(value)


def code_pair(doc, label, before, after):
    doc.add_heading(label, level=3)
    for title, snippet in (("Sebelum", before), ("Sesudah", after)):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        lead = p.add_run(title + ": ")
        lead.bold = True
        run = p.add_run(snippet)
        run.font.name = "Consolas"
        run.font.size = Pt(8)


def add_evidence_image(doc, path, caption):
    if not path.is_file():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(15.5))
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.italic = True
        run.font.size = Pt(8)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.1)
    sec.right_margin = Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in [("Title", 23, DARK), ("Heading 1", 16, GREEN),
                              ("Heading 2", 12, DARK), ("Heading 3", 10, GREEN)]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("TEMU UIN MALANG | TAHAP 4–5").font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Laporan kemajuan proyek | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LAPORAN TAHAP 4 DAN TAHAP 5")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Integrasi, Pengujian, Penyempurnaan, Finalisasi, dan Presentasi")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(GREEN)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TEMU — Lost & Found UIN Malang")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string("C9A227")
    table(doc, ["Identitas", "Keterangan"], [
        ("Nama mahasiswa", "[Isi nama mahasiswa]"),
        ("NIM / kelas", "[Isi NIM dan kelas]"),
        ("Dosen pengampu", "[Isi nama dosen]"),
        ("Tanggal pemeriksaan", "15 Oktober 2026"),
        ("Status", "Laporan kemajuan; sebagian pengujian dan luaran akhir masih tertunda"),
    ])
    doc.add_paragraph(
        "Catatan kejujuran pelaporan: status LULUS hanya diberikan untuk pemeriksaan "
        "yang benar-benar telah dijalankan. Pengujian yang tertahan, umpan balik pengguna, "
        "dan APK yang belum tersedia ditandai secara terbuka; data pengguna uji tidak "
        "dibuat-buat."
    )
    doc.add_page_break()

    doc.add_heading("Ringkasan", level=1)
    doc.add_paragraph(
        "Tahap 4 berfokus pada pemeriksaan integrasi dan mutu fungsi yang telah dirancang "
        "pada tahap sebelumnya, pencatatan masalah, perbaikan, serta uji calon pengguna. "
        "Tahap 5 berfokus pada penyiapan versi final, APK, dokumentasi, repository, video, "
        "presentasi, dan demonstrasi langsung."
    )
    table(doc, ["Aspek", "Kondisi pada 15 Oktober 2026"], [
        ("Flutter analyze", "LULUS — No issues found (pemeriksaan lokal)."),
        ("PHP syntax", "LULUS — config.php dan public/api/index.php tidak memiliki syntax error."),
        ("Flutter tests", "TERTUNDA — proses terhenti karena ruang temporary drive C: tidak cukup; hasil test tidak dapat dinyatakan lulus."),
        ("Uji calon pengguna", "BELUM DILAKSANAKAN / belum ada data umpan balik yang tervalidasi."),
        ("APK kandidat", "BELUM TERSEDIA — belum ditemukan di build/app/outputs/flutter-apk."),
        ("Source control", "Repository remote ada; perubahan lokal masih perlu ditinjau, di-commit, dan diunggah."),
    ], size=8)

    doc.add_heading("BAB I — TAHAP 4: INTEGRASI DAN PENGUJIAN", level=1)
    doc.add_heading("1.1 Cakupan integrasi", level=2)
    bullets(doc, [
        "Flutter menghubungkan halaman login, dashboard, jelajah, form laporan, chat, profil, dan dashboard admin.",
        "PHP REST API menjadi penghubung aplikasi dengan MySQL/MariaDB; Hive digunakan untuk cache lokal.",
        "Fitur utama mencakup laporan lintas akun, chat per laporan, notifikasi, unread badge, kehadiran, dan pengelolaan akun oleh admin.",
        "UI web telah disesuaikan agar memakai ruang desktop yang lebar serta menyusut pada ukuran layar lebih kecil.",
    ])
    doc.add_heading("1.2 Hasil pemeriksaan yang telah dijalankan", level=2)
    table(doc, ["Pemeriksaan", "Hasil", "Bukti / catatan"], [
        ("Flutter static analysis", "LULUS", "No issues found."),
        ("PHP lint config", "LULUS", "backend/config.php: no syntax errors."),
        ("PHP lint API", "LULUS", "backend/public/api/index.php: no syntax errors."),
        ("Widget/unit test", "TERTUNDA", "Flutter gagal menyalin output compiler ke temporary folder karena disk penuh; sekitar 30 MB tersisa sesudah percobaan."),
        ("API admin dan data bersama", "PERNAH DICEK", "Pada pemeriksaan lokal sebelumnya, admin login dan admin sync mendapat respons; ulangi pada versi kandidat akhir."),
        ("Build Web", "PERNAH LULUS", "Build Web pernah berhasil; tetap lakukan build ulang setelah perubahan akhir."),
        ("Build APK", "BELUM DICOBA", "Tidak dijalankan karena ruang disk tidak memadai; APK belum tersedia."),
    ], size=7)

    doc.add_heading("1.3 Matriks skenario pengujian", level=2)
    rows = [
        ("T-01", "Login: email kampus valid, non-kampus, format salah", "Valid diterima; format/identitas salah ditolak", "Belum terverifikasi pada run terakhir"),
        ("T-02", "Identitas email/NIM/nama yang sudah terdaftar", "Email tidak dapat dipakai dengan identitas kedua", "Belum terverifikasi pada run terakhir"),
        ("T-03", "Login admin dan token sesi", "Admin masuk; token salah/kedaluwarsa ditolak", "API pernah dicek; ulangi"),
        ("T-04", "Navigasi dashboard → jelajah/lapor/chat/profil", "Halaman dan tombol kembali bekerja", "Sebagian widget test tersedia; run tertunda"),
        ("T-05", "Form laporan kosong dan valid", "Validasi tampil; laporan valid tersimpan", "Sebagian widget test tersedia; run tertunda"),
        ("T-06", "Kamera dan galeri", "Pilihan sumber foto tersedia; unggah sesuai dukungan perangkat", "Widget chooser ada; kamera fisik perlu diuji"),
        ("T-07", "Sinkronisasi laporan antar akun", "Laporan terlihat dari akun lain dan setelah refresh", "Ulangi pada versi kandidat"),
        ("T-08", "Chat dua akun", "Pesan tersimpan dan terlihat di akun penerima", "Pernah dicek; ulangi dan simpan bukti"),
        ("T-09", "Unread, popup, dan notifikasi", "Jumlah berubah; notifikasi terkait tercatat", "Perlu uji ulang end-to-end"),
        ("T-10", "Online dan last seen", "Online saat heartbeat aktif; waktu terakhir valid saat offline", "Perlu uji lintas dua sesi/perangkat"),
        ("T-11", "Dashboard admin dan detail laporan", "Jumlah total/selesai benar; detail laporan terbuka", "Perlu uji manual dan screenshot"),
        ("T-12", "Kelola akun admin", "Edit tervalidasi; hapus terkonfirmasi; laporan tetap terjaga", "Belum diuji; gunakan akun uji, jangan akun nyata"),
        ("T-13", "Kegagalan API/database", "Pesan error jelas; aplikasi tidak menampilkan data akun lain", "Belum ada bukti uji kondisi gagal"),
        ("T-14", "Responsivitas desktop/tablet/ponsel", "Tidak terpotong, overflow, atau menyisakan area salah", "Perlu uji visual pada beberapa ukuran"),
        ("T-15", "APK Android kandidat", "APK terpasang, API dapat dijangkau, kamera/galeri berfungsi", "Belum dibangun"),
    ]
    table(doc, ["ID", "Skenario", "Hasil yang diharapkan", "Status saat ini"], rows, size=6.5)

    doc.add_heading("1.4 Inventaris test otomatis", level=2)
    bullets(doc, [
        "test/auth_and_navigation_test.dart: validasi AuthService, penolakan email non-kampus, navigasi login valid, dan mode admin.",
        "test/report_and_chat_test.dart: validasi dan submit laporan, pilihan kamera/galeri, badge/status chat, dan kirim pesan.",
        "test/widget_test.dart: smoke test aplikasi.",
        "Jumlah kasus yang terdefinisi: 10. Percobaan eksekusi terbaru terhambat oleh kapasitas disk; setelah ruang tersedia, jalankan flutter test --no-pub dan lampirkan output asli.",
    ])

    doc.add_heading("1.5 Daftar masalah dan perbaikan", level=2)
    table(doc, ["ID", "Temuan", "Tindakan / status"], [
        ("BUG-01", "Permintaan admin di browser gagal preflight karena header X-Admin-Token belum diizinkan CORS.", "DIPERBAIKI: header ditambahkan pada respons CORS; ulangi dari Chrome sebagai verifikasi."),
        ("BUG-02", "Halaman login hanya setinggi konten sehingga sisa viewport tampil putih.", "DIPERBAIKI: kontainer login dibuat memenuhi tinggi viewport."),
        ("BUG-03", "Batas lebar konten membuat tampilan monitor lebar terlihat terlalu kecil.", "DIPERBAIKI: batas lebar responsif dilonggarkan; perlu uji visual ulang desktop/ponsel."),
        ("BUG-04", "Manifest utama Android belum menyatakan izin INTERNET untuk release.", "DIPERBAIKI DI SOURCE: izin ditambahkan ke AndroidManifest.xml; menunggu build APK untuk verifikasi."),
        ("ENV-01", "Ruang drive C: tidak cukup untuk compiler test/APK.", "TERBUKA: kosongkan ruang dengan aman; test terhenti. Tidak ada file yang dihapus oleh pemeriksaan ini."),
        ("ENV-02", "APK kandidat belum tersedia.", "TERBUKA: build dan instal APK setelah disk cukup, izin jaringan, dan URL backend diuji."),
        ("UX-01", "Belum ada catatan uji dari calon pengguna.", "TERBUKA: lakukan sesi nyata, catat temuan dan persetujuan; jangan mengisi hasil fiktif."),
        ("DOC-01", "Perubahan source lokal belum tersimpan di commit remote terbaru.", "TERBUKA: review diff, commit, push, lalu catat tautan commit final."),
    ], size=7)

    doc.add_heading("1.6 Formulir uji calon pengguna", level=2)
    doc.add_paragraph(
        "Isi tabel ini setelah mengamati calon pengguna menjalankan tugas. Gunakan kode "
        "anonim (U1, U2, U3), minta izin sebelum mencatat, dan jangan menulis data pribadi "
        "yang tidak diperlukan. Skala kemudahan: 1 sangat sulit — 5 sangat mudah."
    )
    table(doc, ["Kode / tanggal", "Tugas dan hasil", "Kesulitan / waktu", "Skor 1–5", "Saran pengguna"], [
        ("U1 / ______", "Login; cari laporan; buat laporan; kirim chat", "________________", "___", "________________"),
        ("U2 / ______", "Login; cari laporan; buat laporan; kirim chat", "________________", "___", "________________"),
        ("U3 / ______", "Login; cari laporan; buat laporan; kirim chat", "________________", "___", "________________"),
    ], size=6.5)
    doc.add_paragraph("Kesimpulan dari hasil nyata: __________________________________________________________")
    doc.add_paragraph("Perbaikan UX berdasarkan umpan balik: _________________________________________________")

    doc.add_heading("BAB II — TAHAP 5: FINALISASI DAN PRESENTASI", level=1)
    doc.add_heading("2.1 Daftar kesiapan luaran", level=2)
    table(doc, ["Luaran", "Status", "Langkah penyelesaian"], [
        ("Aplikasi final", "BELUM FINAL", "Selesaikan test, bug kritis, dan uji lintas perangkat."),
        ("APK", "BELUM ADA", "Build release setelah ruang disk cukup; uji instalasi dan koneksi backend."),
        ("Kode sumber", "ADA LOKAL", "Review seluruh perubahan, commit, push, dan pastikan repository dapat dibuka dosen."),
        ("Dokumentasi teknis", "SEBAGIAN ADA", "README/backend README tersedia; perbarui status APK, test final, instalasi, dan keterbatasan."),
        ("Laporan Proyek Akhir", "PERLU DILENGKAPI", "Gabungkan hasil Tahap 1–5 dan isi bukti uji/pengguna yang sesungguhnya."),
        ("Video demonstrasi", "BELUM ADA", "Rekam alur masalah → kebutuhan → solusi → fitur → hasil test."),
        ("Bahan presentasi", "BELUM ADA", "Siapkan slide ringkas sesuai alur penilaian dosen."),
        ("Demo langsung", "DAPAT DISIAPKAN", "Laragon MySQL aktif; buka project di VS Code; pilih TEMU - Demo Web di Chrome; tekan F5."),
    ], size=7)

    doc.add_heading("2.2 Panduan build APK kandidat", level=2)
    doc.add_paragraph(
        "Pastikan Android SDK/Java tersedia, drive memiliki ruang kosong yang cukup, "
        "izin internet sudah masuk ke manifest utama, dan API memakai alamat yang dapat "
        "dijangkau perangkat. Untuk emulator gunakan konfigurasi API emulator; untuk ponsel "
        "fisik gunakan alamat server pada jaringan atau layanan HTTPS. Jangan mengandalkan "
        "127.0.0.1 pada APK ponsel karena alamat itu menunjuk ke ponsel itu sendiri."
    )
    p = doc.add_paragraph()
    run = p.add_run("flutter build apk --release --no-pub")
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    doc.add_paragraph(
        "Lokasi keluaran yang diharapkan: build/app/outputs/flutter-apk/app-release.apk. "
        "Konfigurasi saat ini menggunakan debug signing untuk build release; itu dapat "
        "berguna untuk kandidat uji internal, tetapi bukan konfigurasi distribusi resmi."
    )

    doc.add_heading("2.3 Urutan demonstrasi dan presentasi", level=2)
    bullets(doc, [
        "Masalah: laporan barang hilang/temuan tersebar dan sulit dilacak.",
        "Kebutuhan dan rancangan: aktor mahasiswa/admin, alur, navigasi, serta model data.",
        "Implementasi: Flutter, REST API PHP, MySQL/MariaDB, cache lokal, dan alasan pemilihan.",
        "Demo: login, dashboard, jelajah, buat laporan, akun kedua, chat/notifikasi, lalu dashboard admin.",
        "Kualitas: tunjukkan tabel test, hasil nyata, masalah yang ditemukan, dan bukti perbaikannya.",
        "Keterbatasan dan rencana: jelaskan secara transparan polling chat, kebutuhan backend aktif, dan status APK.",
    ])

    doc.add_heading("2.4 Transparansi penggunaan alat bantu", level=2)
    doc.add_paragraph(
        "Contoh pernyataan yang perlu disesuaikan mahasiswa: ‘Dalam proses pengembangan, "
        "saya menggunakan OpenAI ChatGPT/Codex sebagai alat bantu untuk diskusi debugging, "
        "penyempurnaan antarmuka, dan penyusunan dokumentasi. Saya meninjau perubahan, "
        "menjalankan pemeriksaan yang tercatat, dan bertanggung jawab memahami serta "
        "menjelaskan kode yang digunakan.’ Cantumkan pustaka/API lain sesuai penggunaan "
        "sebenarnya dan kebijakan dosen."
    )

    doc.add_heading("Kesimpulan status", level=1)
    doc.add_paragraph(
        "Fitur inti TEMU dan beberapa pemeriksaan teknis sudah tersedia. Namun Tahap 4 "
        "belum dapat dinyatakan selesai karena test otomatis belum memperoleh hasil lulus, "
        "uji calon pengguna belum terdokumentasi, dan sejumlah skenario integrasi perlu "
        "diulang pada kandidat akhir. Tahap 5 belum selesai karena APK, video, slide, "
        "laporan final gabungan, serta commit/push source terbaru masih perlu disiapkan."
    )

    doc.add_page_break()
    doc.add_heading("PEMBARUAN HASIL PENGUJIAN DAN BUKTI PERBAIKAN", level=1)
    doc.add_paragraph(
        "Bagian ini merupakan pembaruan yang lebih baru dan menggantikan status uji "
        "sementara pada halaman sebelumnya. Pada 15 Oktober 2026, flutter test --no-pub "
        "berhasil dengan hasil All tests passed (10 test). Build APK sudah dicoba ulang "
        "setelah ruang disk tersedia, tetapi gagal sebelum kompilasi aplikasi karena "
        "Android SDK tidak menemukan/memasang NDK 28.2.13676358 (Package ndk not found). "
        "Karena itu APK belum tersedia; ini hambatan toolchain, bukan hasil test Dart."
    )
    doc.add_heading("Perubahan kode sebelum dan sesudah", level=2)
    code_pair(
        doc,
        "BUG-01 — CORS admin",
        "Access-Control-Allow-Headers: Content-Type, Accept",
        "Access-Control-Allow-Headers: Content-Type, Accept, X-Admin-Token",
    )
    code_pair(
        doc,
        "BUG-02 — Login mengisi viewport",
        "body: Container(\n  decoration: const BoxDecoration(...)",
        "body: Container(\n  width: double.infinity,\n  height: double.infinity,\n  decoration: const BoxDecoration(...)",
    )
    code_pair(
        doc,
        "BUG-03 — Layout desktop responsif",
        "ResponsivePage default maxWidth: 1120",
        "ResponsivePage default maxWidth: 1800; dashboard maxWidth: 1650",
    )
    code_pair(
        doc,
        "BUG-04 — Internet Android",
        "Manifest belum mendeklarasikan izin INTERNET",
        '<uses-permission android:name="android.permission.INTERNET" />',
    )
    code_pair(
        doc,
        "Admin — data pengguna",
        "Snapshot admin belum memiliki daftar pengguna",
        "final List<AdminUser> users; // NIM, email, nama, dibuat, online, last seen",
    )
    add_evidence_image(
        doc,
        SCREENSHOT_ADMIN_BEFORE,
        "Bukti 1. Kondisi dashboard admin sebelum perbaikan, screenshot pengguna 9 Oktober 2026.",
    )
    add_evidence_image(
        doc,
        SCREENSHOT_LOGIN_BEFORE,
        "Bukti 2. Halaman web menyisakan ruang putih sebelum perbaikan viewport, screenshot pengguna 9 Oktober 2026.",
    )
    doc.add_paragraph(
        "Gambar yang tersedia adalah bukti sebelum perbaikan. Screenshot sesudah belum "
        "tersedia/terverifikasi sehingga tidak direkayasa; ambil screenshot terbaru saat demo."
    )
    doc.add_heading("Cara melaksanakan uji calon pengguna pada localhost", level=2)
    bullets(doc, [
        "Untuk sesi langsung di komputer pemilik, jalankan backend dan Flutter web; calon pengguna memakai Chrome pada komputer itu. Buat sesi terpisah dengan profil Chrome berbeda atau mode privat agar akun tidak tertukar.",
        "Untuk ponsel di Wi-Fi yang sama, bind backend dan web server ke 0.0.0.0, gunakan alamat IPv4 LAN komputer sebagai API_BASE_URL, izinkan port yang diperlukan di Windows Firewall, lalu buka alamat LAN dari ponsel. Jangan gunakan 127.0.0.1 pada ponsel karena menunjuk ke ponsel itu sendiri.",
        "Untuk peserta dari luar jaringan, deploy sementara ke hosting HTTPS. Jangan membuka port database publik; hanya expose aplikasi/API dan gunakan data akun uji.",
        "Minta persetujuan peserta, gunakan kode anonim U1/U2/U3, amati tugas tanpa mengarahkan, catat waktu/kesulitan/saran, lalu isi formulir. Jangan mengisi hasil sebelum sesi benar-benar dilakukan.",
    ])
    doc.add_paragraph(
        "Uji otomatis lulus tidak menggantikan uji pengguna maupun uji integrasi perangkat. "
        "Tahap 4/5 belum final sampai uji pengguna nyata, skenario integrasi, APK, dan materi presentasi diselesaikan."
    )

    doc.core_properties.title = "Laporan Tahap 4 dan 5 TEMU UIN Malang"
    doc.core_properties.subject = "Integrasi, pengujian, finalisasi, dan presentasi"
    doc.core_properties.author = "[Isi nama mahasiswa]"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
